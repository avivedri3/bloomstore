#!/usr/bin/env python3
"""Build docs/project-book.docx from the Markdown book.

Design decisions live in docs/docx-style-guide.md. The STYLE values below
must change together with that file. This script does not replace the
Markdown source or the HTML conversion.
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt, RGBColor
except ModuleNotFoundError as exc:
    if exc.name == "docx":
        raise SystemExit(
            "python-docx is required. Install it with: python3 -m pip install -r scripts/requirements-docx.txt"
        ) from exc
    raise

ROOT = Path(__file__).resolve().parents[1]

# Specified in docs/docx-style-guide.md. Change both together.
STYLE = {
    "page_width_mm": 210,
    "page_height_mm": 297,
    "margin_right_cm": 3.0,
    "margin_left_cm": 2.5,
    "margin_top_cm": 2.5,
    "margin_bottom_cm": 2.5,
    "hebrew_font": "Arial",
    "latin_font": "Times New Roman",
    "code_font": "Courier New",
    "body_pt": 12,
    "body_line": 1.5,
    "body_after_pt": 8,
    "chapter_pt": 18,
    "section_pt": 14,
    "caption_pt": 11,
    "code_pt": 10,
    "inline_code_pt": 11,
    "header_pt": 10,
    "toc_section_pt": 12,
    "ink": "1F2933",
    "leaf": "3D6B4F",
    "leaf_dark": "2D5040",
    "muted": "5F6B7A",
    "cream_dark": "F0E6DA",
    "code_fill": "F7F4F0",
}

CAPTION_RE = re.compile(r"^(טבלה|איור) \d+ — ")
SECTION_RE = re.compile(r"^\d+\.\d+")
SEPARATOR_CELL_RE = re.compile(r":?-{3,}:?")
FENCE_RE = re.compile(r"^```([\w-]*)\s*$")
ANCHOR_RE = re.compile(r'^<a id="[^"]+"></a>\s*$')
TOC_LINK_RE = re.compile(r"^-?\s*\[.+\]\(#[^)]+\)\s*$")
INLINE_RE = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def mark_bidi(target, enabled: bool = True) -> None:
    element = target._p if hasattr(target, "_p") else target.element
    p_pr = element.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if enabled else "0")


def set_logical_indent(paragraph, start: int, hanging: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:start"), str(start))
    if hanging:
        ind.set(qn("w:hanging"), str(hanging))


def set_paragraph_border(paragraph, edge: str, color: str, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:space"), "1")
    element.set(qn("w:color"), color)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def apply_run_font(run, font: str, size_pt: float, *, bold: bool = False, color: str | None = None, rtl: bool | None = None, fill: str | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = font
    if color:
        run.font.color.rgb = rgb(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    r_fonts.set(qn("w:cs"), font)
    r_fonts.set(qn("w:eastAsia"), font)
    size_half_points = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set(qn("w:val"), size_half_points)
    if rtl is not None:
        rtl_node = r_pr.find(qn("w:rtl"))
        if rtl_node is None:
            rtl_node = OxmlElement("w:rtl")
            r_pr.append(rtl_node)
        rtl_node.set(qn("w:val"), "1" if rtl else "0")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "he-IL" if rtl else "en-US")
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        r_pr.append(shd)


# An English term inside a Hebrew sentence: words, paths, versions, and the
# parentheses that belong to that term only. Sentence punctuation stays outside.
LATIN_PHRASE = re.compile(
    r"\(?@?[A-Za-z][A-Za-z0-9_+#@]*"
    r"(?:[./:\\-][A-Za-z0-9_+#@]+|\s+\d+\+?|\s+[A-Za-z][A-Za-z0-9_+#@]*|\s*[+·]\s*[A-Za-z0-9_+#@]+)*"
    r"\)?"
)
LTR_ISOLATE = "\u2066"
POP_ISOLATE = "\u2069"


def split_scripts(text: str) -> list[tuple[str, str]]:
    parts: list[tuple[str, str]] = []
    position = 0
    for match in LATIN_PHRASE.finditer(text):
        if match.start() > position:
            parts.append(("he", text[position : match.start()]))
        parts.append(("lat", match.group()))
        position = match.end()
    if position < len(text):
        parts.append(("he", text[position:]))
    return parts


def add_text(paragraph, text: str, *, bold: bool = False, size: float | None = None, color: str | None = None) -> None:
    point_size = STYLE["body_pt"] if size is None else size
    ink = color or STYLE["ink"]
    for script, chunk in split_scripts(text):
        if script == "lat":
            run = paragraph.add_run(f"{LTR_ISOLATE}{chunk}{POP_ISOLATE}")
            apply_run_font(run, STYLE["latin_font"], point_size, bold=bold, color=ink, rtl=False)
        else:
            run = paragraph.add_run(chunk)
            apply_run_font(run, STYLE["hebrew_font"], point_size, bold=bold, color=ink, rtl=True)


def add_inline(paragraph, text: str, *, bold: bool = False, size: float | None = None, color: str | None = None) -> None:
    point_size = STYLE["body_pt"] if size is None else size
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            add_text(paragraph, text[position : match.start()], bold=bold, size=point_size, color=color)
        if match.group(1) is not None:
            add_text(paragraph, match.group(1), bold=True, size=point_size, color=color)
        else:
            run = paragraph.add_run(f"{LTR_ISOLATE}{match.group(2)}{POP_ISOLATE}")
            apply_run_font(
                run,
                STYLE["code_font"],
                STYLE["inline_code_pt"] if size is None else min(size, STYLE["inline_code_pt"]),
                bold=bold,
                color=STYLE["ink"],
                rtl=False,
                fill=STYLE["cream_dark"],
            )
        position = match.end()
    if position < len(text):
        add_text(paragraph, text[position:], bold=bold, size=point_size, color=color)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = STYLE["latin_font"]
    normal.font.size = Pt(STYLE["body_pt"])
    normal.font.color.rgb = rgb(STYLE["ink"])
    normal.paragraph_format.line_spacing = STYLE["body_line"]
    normal.paragraph_format.space_after = Pt(STYLE["body_after_pt"])
    normal.paragraph_format.space_before = Pt(0)
    mark_bidi(normal)
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), STYLE["latin_font"])
    r_fonts.set(qn("w:hAnsi"), STYLE["latin_font"])
    r_fonts.set(qn("w:cs"), STYLE["hebrew_font"])
    r_fonts.set(qn("w:eastAsia"), STYLE["hebrew_font"])

    for style_name, size, color in (
        ("Heading 1", STYLE["chapter_pt"], STYLE["ink"]),
        ("Heading 2", STYLE["section_pt"], STYLE["leaf_dark"]),
    ):
        style = document.styles[style_name]
        style.font.name = STYLE["hebrew_font"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.font.italic = False
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(0 if style_name == "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(12 if style_name == "Heading 1" else 6)
        mark_bidi(style)
        heading_r_pr = style.element.get_or_add_rPr()
        heading_fonts = heading_r_pr.find(qn("w:rFonts"))
        if heading_fonts is None:
            heading_fonts = OxmlElement("w:rFonts")
            heading_r_pr.append(heading_fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            heading_fonts.set(qn(attr), STYLE["hebrew_font"])


def apply_page_geometry(section) -> None:
    section.page_width = Mm(STYLE["page_width_mm"])
    section.page_height = Mm(STYLE["page_height_mm"])
    section.right_margin = Cm(STYLE["margin_right_cm"])
    section.left_margin = Cm(STYLE["margin_left_cm"])
    section.top_margin = Cm(STYLE["margin_top_cm"])
    section.bottom_margin = Cm(STYLE["margin_bottom_cm"])
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def set_section_valign(section, value: str | None) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:vAlign"))
    if existing is not None:
        sect_pr.remove(existing)
    if value:
        node = OxmlElement("w:vAlign")
        node.set(qn("w:val"), value)
        sect_pr.append(node)


def set_page_start(section, start: int) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    node = OxmlElement("w:pgNumType")
    node.set(qn("w:start"), str(start))
    sect_pr.append(node)


def enable_update_fields(document: Document) -> None:
    settings = document.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")
    lang = settings.find(qn("w:themeFontLang"))
    if lang is None:
        lang = OxmlElement("w:themeFontLang")
        settings.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:bidi"), "he-IL")


def add_page_number(paragraph) -> None:
    begin = paragraph.add_run()
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin._r.append(begin_char)

    instruction = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instruction._r.append(instr)

    separate = paragraph.add_run()
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate._r.append(separate_char)

    shown = paragraph.add_run("1")
    apply_run_font(shown, STYLE["hebrew_font"], STYLE["body_pt"], color=STYLE["ink"], rtl=True)

    end = paragraph.add_run()
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end._r.append(end_char)


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark_bidi(header)
    header.paragraph_format.space_after = Pt(4)
    add_text(header, "ספר פרויקט — BloomStore", size=STYLE["header_pt"], color=STYLE["muted"])
    set_paragraph_border(header, "bottom", STYLE["leaf"], "8")

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark_bidi(footer)
    add_page_number(footer)


def classify_cover(line: str) -> tuple[str, str]:
    raw = line.strip()
    text = re.sub(r"^#{1,6}\s*", "", raw)
    if raw.startswith("# ") and "BloomStore" in text:
        return "hero", text
    if raw.startswith("# ") and text == "ספר פרויקט":
        return "doctype", text
    if raw.startswith("## "):
        return "department", text
    if raw.startswith("# "):
        return "institution", text
    if text in {"מוגש על ידי", "בהנחיית"}:
        return "label", text
    if "Full-Stack" in text or "מקוונת" in text:
        return "subtitle", text
    if re.fullmatch(r"[\d\-]+", text):
        return "phone", text
    if "ת״ז" in text or 'ת"ז' in text:
        return "student", text
    if re.search(r"\d{4}", text):
        return "date", text
    if re.search(r"\d", text):
        return "address", text
    return "person", text


COVER_FORMAT = {
    "institution": {"size": 22, "bold": True, "before": 0, "after": 6, "color": "ink", "border": True},
    "department": {"size": 14, "bold": False, "before": 0, "after": 18, "color": "ink", "border": False},
    "doctype": {"size": 18, "bold": True, "before": 8, "after": 4, "color": "ink", "border": False},
    "hero": {"size": 28, "bold": True, "before": 4, "after": 4, "color": "leaf", "border": False},
    "subtitle": {"size": 13, "bold": False, "before": 0, "after": 20, "color": "muted", "border": False},
    "label": {"size": 12, "bold": False, "before": 14, "after": 2, "color": "muted", "border": False},
    "student": {"size": 16, "bold": True, "before": 0, "after": 2, "color": "ink", "border": False},
    "person": {"size": 16, "bold": True, "before": 0, "after": 2, "color": "ink", "border": False},
    "phone": {"size": 12, "bold": False, "before": 0, "after": 2, "color": "ink", "border": False},
    "address": {"size": 12, "bold": False, "before": 0, "after": 2, "color": "ink", "border": False},
    "date": {"size": 12, "bold": False, "before": 16, "after": 0, "color": "ink", "border": False},
}


def render_cover(document: Document, lines: list[str]) -> str:
    author = "אביב לייסטן"
    for line in lines:
        role, text = classify_cover(line)
        if role == "student":
            author = text.split("·")[0].strip()
        spec = COVER_FORMAT[role]
        paragraph = document.add_paragraph()
        paragraph.style = document.styles["Normal"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mark_bidi(paragraph)
        paragraph.paragraph_format.space_before = Pt(spec["before"])
        paragraph.paragraph_format.space_after = Pt(spec["after"])
        paragraph.paragraph_format.line_spacing = 1.15
        add_inline(paragraph, text, bold=spec["bold"], size=spec["size"], color=STYLE[spec["color"]])
        if spec["border"]:
            set_paragraph_border(paragraph, "bottom", STYLE["leaf"], "12")
    return author


def render_toc(document: Document, blocks: list[tuple]) -> None:
    title = document.add_paragraph()
    title.style = document.styles["Normal"]
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark_bidi(title)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.keep_with_next = True
    add_text(title, "תוכן עניינים", bold=True, size=STYLE["chapter_pt"], color=STYLE["ink"])
    set_paragraph_border(title, "bottom", STYLE["leaf"], "8")

    for kind, payload in blocks:
        if kind != "heading":
            continue
        level, text = payload
        include = (level == 1 and text.startswith("פרק")) or (
            level == 2 and (SECTION_RE.match(text) is not None or text == "מקורות")
        )
        if not include:
            continue
        paragraph = document.add_paragraph()
        paragraph.style = document.styles["Normal"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        mark_bidi(paragraph)
        paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 1)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.15
        if level == 2:
            set_logical_indent(paragraph, 420)
        add_inline(
            paragraph,
            text,
            bold=(level == 1),
            size=STYLE["section_pt"] if level == 1 else STYLE["toc_section_pt"],
            color=STYLE["ink"] if level == 1 else STYLE["leaf_dark"],
        )


def add_body_paragraph(document: Document):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    mark_bidi(paragraph)
    paragraph.paragraph_format.line_spacing = STYLE["body_line"]
    paragraph.paragraph_format.space_after = Pt(STYLE["body_after_pt"])
    paragraph.paragraph_format.space_before = Pt(0)
    return paragraph


def render_heading(document: Document, level: int, text: str) -> None:
    paragraph = document.add_paragraph()
    style_name = "Heading 1" if level == 1 else "Heading 2"
    paragraph.style = document.styles[style_name]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark_bidi(paragraph)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = 1.15
    if level == 1 and text.startswith("פרק"):
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(12)
        add_inline(paragraph, text, bold=True, size=STYLE["chapter_pt"], color=STYLE["ink"])
    else:
        paragraph.paragraph_format.page_break_before = False
        paragraph.paragraph_format.space_before = Pt(14 if level == 2 else 0)
        paragraph.paragraph_format.space_after = Pt(6)
        add_inline(paragraph, text, bold=True, size=STYLE["section_pt"], color=STYLE["leaf_dark"])


def render_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark_bidi(paragraph)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = True
    add_inline(paragraph, text, bold=True, size=STYLE["caption_pt"], color=STYLE["ink"])


def render_list_item(document: Document, text: str, ordered: bool) -> None:
    paragraph = add_body_paragraph(document)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(4)
    set_logical_indent(paragraph, 420, 420)
    prefix = "" if ordered else "• "
    add_inline(paragraph, prefix + text)


def render_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    mark_bidi(paragraph, False)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_together = True
    shade_paragraph(paragraph, STYLE["code_fill"])
    set_paragraph_border(paragraph, "left", STYLE["leaf"], "16")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line if line else " ")
        apply_run_font(run, STYLE["code_font"], STYLE["code_pt"], color=STYLE["ink"], rtl=False)


def split_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [cell.strip() for cell in raw.split("|")]


def is_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(SEPARATOR_CELL_RE.fullmatch(cell.replace(" ", "")) for cell in cells)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_cell(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), STYLE["ink"])
        borders.append(element)
    tc_pr.append(borders)


def set_table_rtl(table) -> None:
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "autofit")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    tr_pr.append(marker)


def render_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=columns)
    table.style = "Table Grid"
    table.autofit = False
    set_table_rtl(table)
    font_size = 10 if columns >= 4 else 11
    for row_index, values in enumerate(normalized):
        row = table.rows[row_index]
        if row_index == 0:
            mark_header_row(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if row_index == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            border_cell(cell)
            if row_index == 0:
                shade_cell(cell, STYLE["cream_dark"])
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            mark_bidi(paragraph)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline(paragraph, value, bold=(row_index == 0), size=font_size)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0


def parse_blocks(text: str) -> list[tuple]:
    lines = text.splitlines()
    blocks: list[tuple] = []
    index = 0
    in_cover = False
    cover: list[str] = []
    skipping_toc = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped == '<div dir="rtl">':
            in_cover = True
            index += 1
            continue
        if stripped == "</div>":
            in_cover = False
            blocks.append(("cover", cover))
            index += 1
            continue
        if in_cover:
            if stripped:
                cover.append(stripped)
            index += 1
            continue
        if stripped == "---" or ANCHOR_RE.match(stripped):
            index += 1
            continue
        if stripped == "## תוכן עניינים":
            blocks.append(("toc", None))
            skipping_toc = True
            index += 1
            continue
        if skipping_toc:
            if stripped == "" or TOC_LINK_RE.match(stripped):
                index += 1
                continue
            skipping_toc = False

        fence = FENCE_RE.match(stripped)
        if fence:
            body: list[str] = []
            index += 1
            while index < len(lines) and not stripped_is_fence_end(lines[index]):
                body.append(lines[index])
                index += 1
            index += 1
            blocks.append(("code", body))
            continue

        if stripped.startswith("|"):
            raw_rows = [stripped]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw_rows.append(lines[index].strip())
                index += 1
            parsed = [split_row(row) for row in raw_rows if not is_separator(row)]
            blocks.append(("table", parsed))
            continue

        heading = re.match(r"^(#{2,3})\s+(.+)$", stripped)
        if heading:
            level = 1 if len(heading.group(1)) == 2 else 2
            blocks.append(("heading", (level, heading.group(2).strip())))
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            blocks.append(("bullet", bullet.group(1).strip()))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            blocks.append(("number", stripped))
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        blocks.append(("paragraph", stripped))
        index += 1

    return blocks


def stripped_is_fence_end(line: str) -> bool:
    return line.strip() == "```"


def render(document: Document, blocks: list[tuple]) -> str:
    author = "אביב לייסטן"
    started_body = False
    for kind, payload in blocks:
        if kind == "cover":
            author = render_cover(document, payload)
            continue
        if not started_body:
            document.add_section(WD_SECTION.NEW_PAGE)
            break_paragraph = document.paragraphs[-1]
            break_paragraph.paragraph_format.space_before = Pt(0)
            break_paragraph.paragraph_format.space_after = Pt(0)
            break_paragraph.paragraph_format.line_spacing = 1.0
            body = document.sections[-1]
            apply_page_geometry(body)
            set_section_valign(body, None)
            set_page_start(body, 1)
            configure_header_footer(body)
            started_body = True
        if kind == "toc":
            render_toc(document, blocks)
        elif kind == "heading":
            render_heading(document, payload[0], payload[1])
        elif kind == "paragraph":
            if CAPTION_RE.match(payload):
                render_caption(document, payload)
            else:
                paragraph = add_body_paragraph(document)
                add_inline(paragraph, payload)
        elif kind == "bullet":
            render_list_item(document, payload, ordered=False)
        elif kind == "number":
            render_list_item(document, payload, ordered=True)
        elif kind == "code":
            render_code(document, payload)
        elif kind == "table":
            render_table(document, payload)
    return author


def assert_book(document: Document, output: Path) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    missing = [f"פרק {number}" for number in range(1, 11) if f"פרק {number}" not in text]
    if missing:
        raise SystemExit(f"Missing chapters: {', '.join(missing)}")
    captions = [paragraph.text for paragraph in document.paragraphs if re.match(r"^טבלה \d+ —", paragraph.text)]
    figures = [paragraph.text for paragraph in document.paragraphs if re.match(r"^איור \d+ —", paragraph.text)]
    if len(captions) != 14:
        raise SystemExit(f"Expected 14 table captions, found {len(captions)}")
    if len(figures) != 3:
        raise SystemExit(f"Expected 3 figure captions, found {len(figures)}")
    if len(document.tables) != 14:
        raise SystemExit(f"Expected 14 tables, found {len(document.tables)}")
    for needle in ("אביב לייסטן", "מור ברגיג", "הצהרת הסטודנט", "docker build", "תוכן עניינים"):
        if needle not in text:
            raise SystemExit(f"Missing expected text: {needle}")
    if "<div" in text or "```" in text:
        raise SystemExit("Markdown markup leaked into the document")
    if not zipfile.is_zipfile(output):
        raise SystemExit("Output is not a docx zip")
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml")
    if "w:bidi".encode() not in xml:
        raise SystemExit("Document is missing right-to-left markers")


def build(source: Path, output: Path) -> None:
    document = Document()
    configure_styles(document)
    cover = document.sections[0]
    apply_page_geometry(cover)
    set_section_valign(cover, "center")
    enable_update_fields(document)
    author = render(document, parse_blocks(source.read_text(encoding="utf-8")))
    document.core_properties.author = author
    document.core_properties.title = "BloomStore — ספר פרויקט"
    document.core_properties.subject = "ספר פרויקט, מכללת אורט סינגאלובסקי, מסלול טכנאי תוכנה"
    document.core_properties.language = "he-IL"
    document.core_properties.category = "ספר פרויקט"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    assert_book(document, output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert the BloomStore project book Markdown into a DOCX file.")
    parser.add_argument("--source", type=Path, default=ROOT / "docs" / "project-book.md")
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "project-book.docx")
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit(f"Source not found: {args.source}")
    build(args.source, args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
